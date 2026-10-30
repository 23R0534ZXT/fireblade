"""End-to-end workload on a real >=3-GiB CSV in the standalone Windows GUI."""
import copy
import csv
import hashlib
import json
from pathlib import Path
import time

from PySide6.QtCore import QModelIndex, QTimer, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QLineEdit, QPlainTextEdit, QPushButton

from fireblade.core.engine import Case
from fireblade.core.errors import OperationCancelled
from fireblade.core.models import Filter, FilterGroup, FormatRule, Summary, View
from fireblade.core.rule_store import RuleBook
from fireblade.ui.advanced_dialogs import FilterDialog, SortDialog
from fireblade.validation.gui_driver import GuiRun
from fireblade.validation.windows_check import windows_memory

BASE=9007199254740993


def run_stress(app,spec_path,report_path):
    spec=json.loads(Path(spec_path).read_text(encoding='utf-8'))
    source=Path(spec['source']);rows=spec['rows']
    if source.stat().st_size<3*1024**3:raise ValueError('The qualification CSV must exceed 3 GiB')
    run=GuiRun(app,report_path,'3-GiB GUI workload')
    run.report['dataset']={k:spec[k] for k in ('bytes','rows','sha256','seed','description')}
    run.report['not_covered']=['Human-operated physical Explorer gestures','UNC shares','Actual investigator data','Long-term Windows 10/11 desktop use']
    full_case=run.out/'grande analyse.timeline'
    error=None

    def expect(condition,message):
        if not condition:raise AssertionError(message)

    def count(view):return run.work(lambda:run.window.case.count(view))

    def reset():
        run.button('Réinitialiser')

    def group(columns):
        run.header(columns[0],'Grouper par cette colonne')
        for column in columns[1:]:run.header(column,'Ajouter au groupement')
        expect(run.window.view.groups==columns,'GUI group order is incorrect')

    def node_values(parent=QModelIndex()):
        model=run.window.model
        return [list(model.node(model.index(i,0,parent)).values) for i in range(model.rowCount(parent))]

    def expand(index):
        tree=run.window.tree
        tree.setCurrentIndex(index)
        QTest.keyClick(tree,Qt.Key_Right)
        run.settle()
        expect(tree.isExpanded(index),'Group did not expand through the keyboard')

    def sort_keys(keys):
        def edit(dialog):
            if not isinstance(dialog,SortDialog):raise AssertionError('Expected sort editor')
            while dialog.table.rowCount():
                dialog.table.setCurrentCell(0,0)
                QTest.mouseClick(next(b for b in dialog.findChildren(QPushButton) if b.text()=='Retirer'),Qt.LeftButton)
            for column,kind,descending in keys:
                QTest.mouseClick(next(b for b in dialog.findChildren(QPushButton) if b.text()=='Ajouter'),Qt.LeftButton)
                i=dialog.table.rowCount()-1
                run.choose(dialog.table.cellWidget(i,0),column)
                run.choose(dialog.table.cellWidget(i,1),dialog.table.cellWidget(i,1).findData(kind))
                run.choose(dialog.table.cellWidget(i,2),int(descending))
            run.accept(dialog)
        run.modal(edit)
        run.menu('Affichage','Tri sur plusieurs colonnes…')

    def select_first_two():
        tree=run.window.tree
        for i,modifier in ((0,Qt.NoModifier),(1,Qt.ControlModifier)):
            index=run.window.model.index(i,0)
            tree.scrollTo(index)
            run.app.processEvents()
            QTest.mouseClick(tree.viewport(),Qt.LeftButton,modifier,pos=tree.visualRect(index).center())
        return [run.window.model.node(index).values[0] for index in tree.selectionModel().selectedRows()]

    try:
        expect(app.platformName()=='windows' and run.report['frozen'],'Actual frozen Windows process required')
        run.check('full_import_of_3_GiB_from_readonly_CSV',lambda:run.import_csv(source,full_case))
        def integrity():
            case=run.window.case
            expect(case.info['rows']==rows,'Lost or duplicated input records')
            expect(case.info['sha256']==spec['sha256'],'Source digest mismatch')
            expect(count(View())==rows,'Wrong full count')
            def samples():
                return case.db.execute('SELECT * FROM events WHERE __row_id IN ('+','.join('?' for _ in spec['samples'])+') ORDER BY __row_id',list(map(int,spec['samples']))).fetchall()
            for record in run.work(samples):expect(list(record[1:])==spec['samples'][str(record[0])],'Source values changed at record '+str(record[0]))
            run.capture('01-import-complete')
            return dict(source_bytes=source.stat().st_size,records=rows,database_bytes=full_case.stat().st_size)
        run.check('independent_row_count_hash_unicode_multiline_and_ID_oracle',integrity)

        for number,scenario in enumerate(spec['groups']):
            def exercise(scenario=scenario,number=number):
                group(scenario['columns'])
                parent=QModelIndex()
                for depth,level in enumerate(scenario['levels']):
                    expected=level['expected']
                    expect(node_values(parent)==expected[:250],f'Group counts differ at {scenario["columns"]}, level {depth}')
                    if len(expected)>250:
                        run.window.tree.setCurrentIndex(parent)
                        run.button('Suivant')
                        expect(node_values(parent)==expected[:1000],'Group pagination differs')
                    if depth+1<len(scenario['levels']):
                        parent=run.window.model.index(0,0,parent);expand(parent)
                run.capture(f'group-{number:02d}')
                return dict(columns=scenario['columns'],levels=len(scenario['levels']))
            run.check('group_by_'+str(number)+'_'+str(scenario['columns']),exercise)

        def filtered_groups():
            run.window.set_view(View(filters=[Filter(3,'eq','user042')]))
            run.settle()
            scenario=spec['filtered_groups']
            group(scenario['columns'])
            parent=QModelIndex()
            for depth,level in enumerate(scenario['levels']):
                expect(node_values(parent)==level['expected'][:250],'Filtered group counts incorrect')
                if depth+1<len(scenario['levels']):
                    parent=run.window.model.index(0,0,parent);expand(parent)
            run.capture('filtered-groups')
        run.check('group_by_two_columns_under_active_filter',filtered_groups)

        def repetitions():
            reset()
            group([1,3,4])
            readings=[]
            for cycle in range(40):
                first=run.window.model.index(0,0)
                expand(first)
                second=run.window.model.index(0,0,first)
                expand(second)
                third=run.window.model.index(0,0,second)
                expand(third)
                expect(run.window.model.rowCount(third)>0,'Empty leaf group')
                run.window.tree.setCurrentIndex(third)
                QTest.keyClick(run.window.tree,Qt.Key_Left)
                run.settle()
                run.window.tree.setCurrentIndex(second)
                QTest.keyClick(run.window.tree,Qt.Key_Left);run.settle()
                run.window.tree.setCurrentIndex(first)
                QTest.keyClick(run.window.tree,Qt.Key_Left);run.settle()
                expect(not run.window.model.node(first).children,'Collapsed branch retained descendants')
                if cycle in (5,15,25,39):readings.append(windows_memory()['working_set_mib'])
            expect(readings[-1]-readings[0]<128,'Resident memory keeps growing through repeated group collapse')
            return dict(cycles=40,working_set_mib=readings)
        run.check('40_cycles_three_level_expand_collapse_and_memory',repetitions)

        def search():
            reset()
            run.text(run.window.search,'Enregistrement 137,')
            QTest.keyClick(run.window.search,Qt.Key_Return);run.settle()
            expect(count(run.window.view)==1,'Global search has wrong result count')
            expect(run.window.model.node(run.window.model.index(0,0)).values[0]==138,'Global search matched wrong source record')
            run.capture('02-search')
            reset()
        run.check('global_search_keyboard_and_reset',search)

        def boolean_editor():
            def fill(dialog):
                if not isinstance(dialog,FilterDialog):raise AssertionError('Expected Boolean editor')
                editor=dialog.editor
                run.choose(editor.tree.itemWidget(editor.root,0),1)  # OR
                for user in ('user042','user043'):
                    QTest.mouseClick(next(b for b in editor.findChildren(QPushButton) if b.text()=='+ Condition'),Qt.LeftButton)
                    item=editor.tree.currentItem()
                    run.choose(editor.tree.itemWidget(item,0),3)
                    operators=editor.tree.itemWidget(item,1)
                    run.choose(operators,operators.findData('eq'))
                    run.text(editor.tree.itemWidget(item,2),user)
                run.capture('03-filter-editor',dialog)
                run.accept(dialog)
            run.modal(fill);run.menu('Filtres','Éditeur ET / OU / NON…')
            expect(count(run.window.view)==spec['reference']['user042_or_user043'],'Boolean editor did not combine OR correctly')
            reset()
        run.check('Boolean_OR_editor_mouse_keyboard_against_Python_oracle',boolean_editor)

        filter_cases=[('exact_user',[Filter(3,'eq','user042')],'user042'),
                      ('user_and_artifact',[Filter(3,'eq','user042'),Filter(1,'eq','MFT')],'user042_MFT'),
                      ('RE2_regex',[Filter(3,'regex','^user0(4[0-9])$')],'user_regex'),
                      ('NULL_user',[Filter(3,'empty')],'empty_user'),
                      ('ISO_date_range',[Filter(0,'after','2026-10-31 00:00:00')],'date_after')]
        for name,filters,key in filter_cases:
            def filtering(filters=filters,key=key):
                run.window.set_view(View(filters=filters));run.settle()
                expect(count(run.window.view)==spec['reference'][key],'Filter count differs from independent reference: '+key)
            run.check('filter_'+name,filtering)
        def invalid_regex():
            failures=[]
            run.window.run(lambda:run.window.case.count(View(filters=[Filter(3,'regex','[')])),lambda _:None,
                           exclusive=True,on_error=failures.append)
            run.settle();expect(bool(failures),'Invalid regex not reported')
            reset();expect(count(View())==rows,'Connection unusable after invalid regex')
        run.check('invalid_regex_and_connection_recovery',invalid_regex)

        def numeric_sort():
            reset();sort_keys([(6,'number',False)])
            expect(run.window.model.node(run.window.model.index(0,0)).values[0]==2,'Large numeric ascending sort is inexact')
            sort_keys([(6,'number',True)])
            last=rows-1
            while last%1999==0 or last%991==0:last-=1
            expect(run.window.model.node(run.window.model.index(0,0)).values[0]==last+1,'Large numeric descending sort is inexact')
            run.capture('04-exact-sort')
        run.check('exact_large_integer_sort_both_directions_via_editor',numeric_sort)

        def deep_sort():
            sort_keys([(1,'text',False),(6,'number',True)])
            def expected(offset):
                matches=[]
                for i in range(rows-1,-1,-1):
                    if i%8==3 and i%1999!=0 and i%991!=0:
                        matches.append(i+1)
                        if len(matches)>=offset+250:break
                return matches[offset:offset+250]
            expect([r[0] for r in node_values()]==expected(0),'Multi-column order incorrect')
            for _ in range(5):run.button('Suivant')
            expect(run.window.model.root.offset==1000,'Page window did not evict old rows')
            expect([r[0] for r in node_values()][:250]==expected(1000),'Deep sort cache returned wrong rows')
            run.button('Précédent')
            expect([r[0] for r in node_values()]==expected(0),'Previous page order incorrect')
            run.capture('05-deep-pagination')
        run.check('multi_column_sort_deep_pagination_cache_and_back',deep_sort)

        def timestamp_sort():
            sort_keys([(0,'timestamp',False)])
            candidates=[i for i in range(0,rows,172800) if i%1979!=0 and i%733!=0]
            expect(run.window.model.node(run.window.model.index(0,0)).values[0]==min(candidates)+1,'ISO timestamp ascending order incorrect')
            sort_keys([(0,'timestamp',True)])
            candidates=[i for i in range(172799,rows,172800) if i%1979!=0 and i%733!=0]
            expect(run.window.model.node(run.window.model.index(0,0)).values[0]==min(candidates)+1,'ISO timestamp descending order incorrect')
        run.check('ISO_timestamp_sort_both_directions',timestamp_sort)

        def colors():
            rule=FormatRule(FilterGroup('and',[Filter(3,'eq','user042')]),'#ffd5cc',target=3)
            run.window.rule_store.replace(RuleBook(spec['headers'],[rule]));run.window.reload_global_rules();run.settle()
            run.window.set_view(View(filters=[Filter(3,'eq','user042')],format_rules=[FormatRule(rule.condition,'#ccddee',target=3)]));run.settle()
            expect(run.window.model.data(run.window.model.index(0,4),Qt.BackgroundRole).name()=='#ccddee','Local format priority incorrect')
            run.capture('06-rules')
        run.check('global_and_local_cell_rules_on_large_filtered_data',colors)

        def notes():
            selected=select_first_two();expect(len(selected)==2,'Mouse selection did not select two rows')
            run.menu('Édition','Marquer la sélection')
            expect(set(run.window.case.bookmarks())==set(selected),'Bulk mark affected wrong rows')
            select_first_two()
            def fill(dialog):
                field=dialog.findChild(QPlainTextEdit)
                if field is None:raise AssertionError('Note text editor not found')
                run.text(field,'À examiner\n=1+1');run.accept(dialog)
            run.modal(fill);run.menu('Édition','Ajouter une note à la sélection…')
            expect(all(v=='À examiner\n=1+1' for v in run.window.case.bookmarks().values()),'Unicode notes altered')
            QTest.mouseClick(run.window.only_notes,Qt.LeftButton);run.settle()
            expect(count(run.window.view)==2,'Bookmarks checkbox did not filter')
            run.capture('07-notes');QTest.mouseClick(run.window.only_notes,Qt.LeftButton);run.settle()
            return dict(source_record_ids=selected)
        run.check('mouse_selection_bulk_notes_and_bookmarks_checkbox',notes)

        def save_and_reopen():
            def fill(dialog):run.text(dialog.findChild(QLineEdit),'triage 3GiB');run.accept(dialog)
            run.modal(fill);run.menu('Analyse','Enregistrer la vue…')
            saved=run.work(run.window.case.views)['triage 3GiB']
            run.window.tree.setColumnHidden(14,True)
            session=run.out/'session.fireblade-session'
            run.window.save_session(session)
            small=run.out/'small.csv';small.write_text('Name\nDeuxième analyse\n',encoding='utf-8')
            run.import_csv(small,run.out/'small.timeline')
            expect(run.window.tabs.count()==2,'Second tab missing')
            # Click the first tab as a user would.
            QTest.mouseClick(run.window.tabs,Qt.LeftButton,pos=run.window.tabs.tabRect(0).center());run.settle()
            expect(run.window.case.path==full_case,'Wrong analysis after tab switch')
            expect(run.window.tree.isColumnHidden(14),'Column layout lost')
            run.window.load_session(session);run.settle()
            expect(run.window.case.path==full_case and run.window.tree.isColumnHidden(14),'Session restore failed')
            expect(len(run.window.case.bookmarks())==2,'Persisted notes lost')
            run.modal(lambda dialog:run.accept(dialog))
            run.menu('Analyse','Charger une vue…')
            expect(run.window.view==saved,'Named view did not restore filters, sort, colors and column layout')
            run.window.tree.setColumnHidden(14,True)
            run.capture('08-session')
        run.check('save_view_switch_tabs_release_database_and_reopen_session',save_and_reopen)

        def filtered_exports():
            case=run.window.case;view=View(filters=[Filter(3,'eq','user042')])
            for extension in ('csv','xlsx'):
                target=run.out/('filtered.'+extension)
                run.work(lambda:case.export(target,view,include_annotations=True),600)
                def verify_filtered():
                    if extension=='csv':
                        with target.open(encoding='utf-8-sig',newline='') as incoming:
                            reader=csv.reader(incoming);next(reader);n=0
                            for record in reader:
                                expect(record[3]=='user042','Export filter incorrect');n+=1
                        return n
                    from openpyxl import load_workbook
                    book=load_workbook(target,read_only=True,data_only=False)
                    try:
                        iterator=book.active.iter_rows();next(iterator)
                        return sum(1 for _ in iterator)
                    finally:book.close()
                n=run.work(verify_filtered,600)
                expect(n==spec['reference']['user042'],'Filtered export row count incorrect')
                target.unlink()
            target=run.out/'grouped.csv'
            run.work(lambda:case.export(target,View(groups=[1]),grouped=True))
            with target.open(encoding='utf-8-sig',newline='') as incoming:
                actual=list(csv.reader(incoming))[1:]
            expected=spec['groups'][0]['levels'][0]['expected']
            expect(actual==[[value,str(n)] for value,n in expected],'Grouped export counts incorrect')
            target.unlink()
        run.check('CSV_XLSX_annotations_and_group_export_full_filtered_results',filtered_exports)

        def replace_remove_notes():
            select_first_two()
            def fill(dialog):
                run.text(dialog.findChild(QPlainTextEdit),'Note remplacée');run.accept(dialog)
            run.modal(fill);run.menu('Édition','Remplacer les notes de la sélection…')
            expect(set(run.window.case.bookmarks().values())=={'Note remplacée'},'Replace notes failed')
            select_first_two();run.menu('Édition','Retirer les signets de la sélection')
            expect(run.window.case.bookmarks()=={},'Remove notes failed')
            QTest.mouseClick(run.window.only_notes,Qt.LeftButton);run.settle()
            expect(count(run.window.view)==0 and run.window.model.rowCount()==0,'Empty bookmark view is incorrect')
            reset();expect(count(View())==rows,'Reset after empty results failed')
        run.check('replace_remove_selected_notes_and_empty_result_recovery',replace_remove_notes)

        def full_export():
            target=run.out/'full-export.csv';case=run.window.case
            run.work(lambda:case.export(target,View()),900)
            def verify():
                total=0
                with target.open(encoding='utf-8-sig',newline='') as incoming:
                    reader=csv.reader(incoming);expect(next(reader)==spec['headers'],'Export header altered')
                    for record in reader:
                        total+=1
                        if str(total) in spec['samples']:
                            expect([v or None for v in record]==spec['samples'][str(total)],'Full export values altered')
                with target.open('rb') as incoming:
                    exported_hash=hashlib.file_digest(incoming,'sha256').hexdigest()
                expect(exported_hash==spec['sha256'],'Full CSV round-trip changed source bytes or record order')
                return total
            expect(run.work(verify,900)==rows,'Full multi-GB export dropped rows')
            size=target.stat().st_size;target.unlink()
            return dict(exported_records=rows,exported_bytes=size)
        run.check('successful_full_multi_GiB_CSV_export_and_independent_readback',full_export)

        def cancelled_export():
            target=run.out/'cancel-export.csv';target.write_bytes(b'previous export')
            cancelled=[]
            run.window.run(lambda:run.window.case.export(target,View()),lambda _:None,exclusive=True,on_cancel=lambda:cancelled.append(True))
            timer=QTimer()
            def cancel():
                partial=target.with_name(target.name+'.partial')
                if partial.exists() and partial.stat().st_size>1024*1024:
                    timer.stop();run.window.interrupt_queries()
            timer.timeout.connect(cancel);timer.start(25)
            run.settle(600);timer.stop()
            expect(cancelled==[True] and target.read_bytes()==b'previous export','Cancelled export replaced prior data')
            expect(not target.with_name(target.name+'.partial').exists(),'Partial export leaked')
            target.unlink();expect(count(View())==rows,'Case unusable after cancellation')
        run.check('cancel_multi_GiB_export_preserve_target_and_reuse_connection',cancelled_export)

        def cancelled_import():
            run.window.set_view(View(filters=[Filter(3,'eq','user042')],groups=[1],sort_column=6,sort_type='number'))
            run.settle();run.window.tree.setColumnHidden(14,True)
            previous=copy.deepcopy(run.window.view);native=Case.import_csv
            audit={};timer=QTimer()
            def traced(case,source,delimiter=None,encoding='utf-8',progress=lambda _:None):
                def forward(message):
                    if message=='Écriture des enregistrements sur disque…':audit['write_started']=time.monotonic()
                    elif message=='Vérification de l’intégrité de la source…':audit['native_finished']=True
                    elif message.startswith('Import sur disque · ') and message.endswith(' enregistrements écrits…'):
                        audit['written_records']=int(message.split(' · ')[1].split(' ')[0].replace(',',''))
                    progress(message)
                try:return native(case,source,delimiter,encoding,forward)
                except OperationCancelled as exc:
                    audit['cancel_exception_cause']=type(exc.__cause__).__name__
                    audit['tables_after_rollback']=case.db.execute('SHOW TABLES').fetchall()
                    raise
            def cancel():
                if audit.get('written_records',0)>0 and not audit.get('native_finished'):
                    timer.stop();run.window.interrupt_queries()
            Case.import_csv=traced;timer.timeout.connect(cancel);timer.start(25)
            try:run.import_csv(source,run.out/'cancelled.timeline')
            finally:timer.stop();Case.import_csv=native
            expect(audit.get('written_records',0)>0,'Import was not interrupted after actual database writes')
            expect(audit.get('tables_after_rollback')==[],'Partial tables survived rollback')
            expect(run.window.case.path==full_case and run.window.view==previous,'Previous analysis/view not restored')
            expect(run.window.tree.isColumnHidden(14),'Column layout not restored after cancelled import')
            return audit
        run.check('cancel_partial_3_GiB_import_rollback_and_restore_prior_workspace',cancelled_import)

        def advanced_groups():
            view=View(groups=[1], summaries=[Summary(11,'sum')],group_order='summary')
            run.window.set_view(view);run.settle()
            actual=node_values()
            expected=sorted(spec['advanced']['size_sums'].items(),key=lambda item:(-item[1],item[0]))
            expect([(row[0],int(row[2])) for row in actual]==expected,'Large-data group sums differ from independent generator')
            view=View(groups=[0],group_intervals={'0':'day'},group_order='value',group_descending=False)
            run.window.set_view(view);run.settle()
            expected=sorted(spec['advanced']['day_counts'],key=lambda item:(item[0] is None,item[0] or ''))
            expect(node_values()==expected,'Large-data temporal groups differ from independent generator')
            run.window.pin_column(1);run.app.processEvents()
            expect(run.window.tree.frozen.columnWidth(2)>0,'Pinned grid is not visible with large data')
            run.capture('10-advanced-groups-pinned')
            run.window.pin_column(1)
        run.check('advanced_summaries_temporal_groups_and_pinned_3_GiB_grid',advanced_groups)

        def projected_export():
            from fireblade.core.filter_text import parse_filter
            view=View(expression=parse_filter("[User] = 'user042'",run.window.case.headers))
            target=run.out/'visible-columns.csv'
            run.work(lambda:run.window.case.export(target,view,columns=[6,3,1]),300)
            with target.open(encoding='utf-8-sig',newline='') as incoming:
                reader=csv.reader(incoming)
                expect(next(reader)==['RecordID','User','Artifact'],'Projection columns are wrong')
                total=0
                for row in reader:
                    expect(row[1]=='user042','Projection contains a filtered-out record')
                    total+=1
            expect(total==spec['reference']['user042'],'Projected export lost filtered records')
            target.unlink()
        run.check('text_expression_and_visible_column_export_on_3_GiB',projected_export)

        def search_large_and_small():
            from fireblade.ui.cross_search import WorkspaceSearch
            from fireblade.core.models import AnalysisSnapshot
            small_source=run.out/'cross-search.csv'
            small_source.write_text('User,Text\nuser042,one\nother,two\n')
            run.window.remember_active()
            current=copy.deepcopy(next(doc for doc in run.window.documents if doc.path==str(full_case)))
            def operation():
                small=Case(run.out/'cross-search.timeline',settings=run.window.settings)
                try:
                    small.import_csv(small_source)
                    snapshot=AnalysisSnapshot(str(small.path),small.info,View(),b'')
                finally:small.close()
                task=WorkspaceSearch([current,snapshot],run.window.settings,run.window.case)
                return task.find('user042')
            hits=run.work(operation,300)
            expect([hit.count for hit in hits]==[spec['reference']['user042'],1],'Cross-file search differs on large and small analyses')
        run.check('cross_file_search_in_full_3_GiB_and_second_analysis',search_large_and_small)

        def final_integrity():
            def digest():
                with source.open('rb') as incoming:return hashlib.file_digest(incoming,'sha256').hexdigest()
            expect(run.work(digest,300)==spec['sha256'],'Source changed during GUI workload')
            expect(count(View())==rows,'Final record count changed')
            memory=windows_memory()
            expect(memory['peak_working_set_mib']<1536,'Process exceeded the QA runaway-memory threshold of 1536 MiB')
            run.capture('09-final-state')
            return dict(source_unchanged=True,final_rows=rows,**memory)
        run.check('final_source_integrity_case_reuse_and_memory',final_integrity)
    except Exception as exc:error=exc
    return run.finish(error)
