"""Native Windows acceptance check, also executed inside the frozen EXE.

No pytest or additional runtime dependency. The report records what was actually
run; Explorer drag/drop and human use of native file pickers remain manual checks.
"""
import csv
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import platform
import stat
import struct
import sys
import tempfile
import time

import duckdb
from PySide6 import __version__ as pyside_version
from PySide6.QtCore import QItemSelectionModel, QTimer, Qt, qVersion
from PySide6.QtWidgets import QDialog

from fireblade.ui.app import Window
from fireblade.ui.dialogs import ImportDialog
from fireblade.core.models import Filter, FilterGroup, FormatRule, SortKey, View
from fireblade.core.rule_store import RuleBook, RuleStore
from fireblade.runtime.windows_paths import analyses_directory
from fireblade.runtime.resources import build_info, default_config, resource_path


def windows_memory():
    import ctypes
    from ctypes import wintypes
    class Counters(ctypes.Structure):
        _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD)] + [
            (name, ctypes.c_size_t) for name in (
                "PeakWorkingSetSize", "WorkingSetSize", "QuotaPeakPagedPoolUsage",
                "QuotaPagedPoolUsage", "QuotaPeakNonPagedPoolUsage", "QuotaNonPagedPoolUsage",
                "PagefileUsage", "PeakPagefileUsage")]
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.GetCurrentProcess.restype = wintypes.HANDLE
    kernel.K32GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
    kernel.K32GetProcessMemoryInfo.restype = wintypes.BOOL
    counters = Counters()
    counters.cb = ctypes.sizeof(counters)
    if not kernel.K32GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
        raise ctypes.WinError(ctypes.get_last_error())
    return dict(peak_working_set_mib=round(counters.PeakWorkingSetSize / 1024**2, 1),
                working_set_mib=round(counters.WorkingSetSize / 1024**2, 1))


def run_check(app, report_path):
    report_path = Path(report_path).resolve()
    if report_path.exists():
        raise ValueError("Le rapport existe déjà ; choisissez un nouveau fichier.")
    report = dict(status="running", started_utc=datetime.now(timezone.utc).isoformat(),
                  platform=sys.platform, operating_system=platform.platform(),
                  machine=platform.machine(), pointer_bits=struct.calcsize("P") * 8,
                  python=sys.version.split()[0], duckdb=duckdb.__version__,
                  qt=qVersion(), pyside=pyside_version,
                  frozen=bool(getattr(sys, "frozen", False)),
                  packaging=build_info().get("packaging"),
                  executable=sys.executable, qt_platform=app.platformName(), checks=[],
                  not_covered=["Explorateur Windows : glisser-déposer réel", "Sélection manuelle dans les dialogues natifs",
                               "CSV forensic réels de plusieurs Go", "Partages UNC", "Essais prolongés Windows 10/11"])
    window, work = None, None
    readonly = []
    errors = []

    def expect(condition, message):
        if not condition:
            raise AssertionError(message)

    def settle(timeout=30):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            app.processEvents()
            if not window.jobs and not window.pool.activeThreadCount() and (window.model is None or not window.model.jobs):
                return
            time.sleep(.005)
        raise TimeoutError("L’interface n’a pas terminé ses opérations en arrière-plan.")

    def check(name, operation):
        started = time.perf_counter()
        try:
            operation()
            expect(not errors, f"Erreurs de l’interface : {errors}")
            report["checks"].append(dict(name=name, status="passed", seconds=round(time.perf_counter() - started, 4)))
        except Exception as exc:
            report["checks"].append(dict(name=name, status="failed", error=str(exc)))
            raise

    def import_with_dialog(source, target, sheet=None):
        def accept_dialog():
            dialog = app.activeModalWidget()
            if not isinstance(dialog, ImportDialog):
                errors.append("La fenêtre de confirmation d’import ne s’est pas ouverte.")
                if isinstance(dialog, QDialog):
                    dialog.reject()
                return
            dialog.target.setText(str(target))
            if sheet is not None:
                dialog.sheet.setCurrentText(sheet)
            dialog.accept()
        QTimer.singleShot(0, accept_dialog)
        window.open_path(source)

    try:
        expect(os.name == "nt" and sys.platform == "win32", "Ce contrôle exige un véritable processus Windows.")
        expect(report["pointer_bits"] == 64, "Un Python ou exécutable 64 bits est requis.")
        expect(app.platformName() == "windows", "Le plugin Qt Windows natif doit être utilisé, pas offscreen.")
        work = tempfile.TemporaryDirectory(prefix="Fireblade vérification ")
        directory = Path(work.name).resolve()  # Windows TEMP may contain an 8.3 alias such as RUNNER~1.
        source = directory / "événements Windows.csv"
        with source.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.writer(output, lineterminator="\r\n")
            writer.writerow(["Artifact", "User", "RecordID", "Timestamp", "Description"])
            writer.writerows([
                ["MFT", "élève", "0001", "2026-10-30 12:00:00", "C:\\Windows\\cmd.exe"],
                ["EVTX", "bob", "9007199254740993", "2026-10-30 12:00:01", "Ligne A\nLigne B"],
                ["MFT", "élève", "2", "", "Donnée accentuée"],
            ])
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        source.chmod(stat.S_IREAD)
        readonly.append(source)
        case_path = directory / "analyse été.timeline"
        window = Window(rule_path=directory / "règles globales.json")
        window.error = lambda message: errors.append(str(message))
        window.show()
        app.processEvents()

        def startup():
            expect(window.isVisible() and int(window.winId()) != 0, "La fenêtre Windows n’est pas visible.")
            expect(window.content_stack.currentIndex() == 0, "L’accueil doit être affiché.")
            expected_directory = analyses_directory()
            dialog = ImportDialog(source, window)
            expect(Path(dialog.target.text()).parent == expected_directory, "La destination doit être dans les données utilisateur.")
            dialog.deleteLater()
        check("native_window_and_user_data_destination", startup)

        def embedded_resources():
            from fireblade.core.settings import load_settings
            expect(load_settings(default_config()).memory_mb == 256, "La configuration intégrée n’a pas été lue.")
            expect("Fireblade" in resource_path("docs/guides/UTILISATION.md").read_text(encoding="utf-8"), "Le guide intégré est absent.")
            expect("MIT License" in resource_path("LICENSE").read_text(encoding="utf-8"), "La licence du projet n’est pas intégrée.")
            if report["frozen"]:
                expect(report["packaging"] == "onefile", "Le livrable doit être un exécutable onefile.")
                dependencies = json.loads(resource_path("legal/DEPENDENCIES.json").read_text(encoding="utf-8"))
                expect({"duckdb", "PySide6-Essentials", "openpyxl"}.issubset({item["name"] for item in dependencies}), "Les notices des dépendances sont absentes.")
                expect(resource_path("legal/qt/LGPL-3.0-only.txt").is_file(), "La licence Qt n’est pas intégrée.")
                for name in ("vcruntime140.dll", "vcruntime140_1.dll", "msvcp140.dll"):
                    expect((Path(sys._MEIPASS) / name).is_file(), f"Le runtime Windows {name} n’est pas intégré.")
                expect(sorted(p.name for p in Path(sys.executable).parent.iterdir()) == [Path(sys.executable).name], "Le programme a besoin de fichiers à côté du .exe.")
        check("single_file_embedded_guide_defaults_and_licenses", embedded_resources)

        def importing():
            import_with_dialog(source, case_path)
            settle()
            expect(window.case is not None and window.case.path == case_path,
                   f"L’analyse n’a pas été ouverte. Erreurs : {errors}")
            rows = window.case.page(View())
            expect(rows[0][2:4] == ("élève", "0001"), "Accents ou zéros initiaux altérés.")
            expect(rows[1][3] == "9007199254740993" and rows[1][-1] == "Ligne A\nLigne B", "Grand entier ou cellule multiligne altéré.")
            expect(hashlib.sha256(source.read_bytes()).hexdigest() == source_hash, "La source en lecture seule a changé.")
        check("readonly_utf8_bom_crlf_unicode_paths_import", importing)

        def grouping_sorting():
            window.group_column(0)
            settle()
            expect(window.case.page(window.view) == [("MFT", 2), ("EVTX", 1)], "Comptages de groupes incorrects.")
            window.group_column(1, append=True)
            settle()
            expect(window.case.page(window.view, ("MFT",)) == [("élève", 2)], "Sous-groupes incorrects.")
            window.set_view(View(sort_column=2, sort_type="number", descending=True))
            settle()
            expect([r[3] for r in window.case.page(window.view)] == ["9007199254740993", "2", "0001"], "Tri numérique inexact.")
            window.filter_value(1, "élève")
            settle()
            expect(window.case.count(window.view) == 2, "Filtre rapide incorrect.")
        check("nested_groups_exact_numeric_sort_and_filter", grouping_sorting)

        target = directory / "export vérifié.csv"
        def locked_export():
            target.write_bytes(b"previous export")
            denied = False
            with target.open("rb"):
                try:
                    window.case.export(target, View())
                except PermissionError:
                    denied = True
            expect(denied, "Le remplacement d’un fichier verrouillé devait être refusé par Windows.")
            expect(target.read_bytes() == b"previous export", "L’ancien export a été altéré.")
            expect(not target.with_name(target.name + ".partial").exists(), "Un export temporaire a été laissé après l’échec.")
            window.case.export(target, View())
            with target.open(encoding="utf-8-sig", newline="") as incoming:
                rows = list(csv.reader(incoming))
            expect(len(rows) == 4 and rows[2][2] == "9007199254740993", "L’export après déverrouillage est incorrect.")
        check("windows_locked_export_and_retry", locked_export)

        def clipboard():
            index = window.model.index(0, 0)
            window.tree.setCurrentIndex(index)
            window.tree.selectionModel().select(index, QItemSelectionModel.ClearAndSelect | QItemSelectionModel.Rows)
            window.copy_rows()
            app.processEvents()
            expect("RecordID" in app.clipboard().text() and "élève" in app.clipboard().text(), "Copie Unicode dans le presse-papiers incorrecte.")
        check("windows_unicode_clipboard", clipboard)

        saved_view = View(filters=[Filter(1, "eq", "élève")], sort_column=2, sort_type="number")
        def reopen():
            window.case.bookmark(1, "À conserver")
            window.case.save_view("triage", saved_view)
            window.detach()
            moved = directory / "analyse déplacée.timeline"
            case_path.rename(moved)
            moved.rename(case_path)  # Proves the closed DB released its Windows file handle.
            window.open_path(case_path)
            settle()
            expect(window.case.bookmarks() == {1: "À conserver"}, "Signet perdu après fermeture.")
            window.set_view(window.case.views()["triage"])
            settle()
            expect(window.view == saved_view and window.case.count(window.view) == 2, "Vue persistée incorrecte.")
        check("released_file_handle_bookmarks_and_saved_view_reopen", reopen)

        def cancelled_import():
            larger = directory / "import à annuler.csv"
            block = (b"0001,alice," + b"artifact detail " * 20 + b"\r\n") * 10000
            with larger.open("wb") as output:
                output.write(b"ID,User,Description\r\n")
                for _ in range(20):
                    output.write(block)
            window.tree.setColumnHidden(5, True)
            import_with_dialog(larger, directory / "annulé.timeline")
            expect(window.busy, "L’import de contrôle s’est terminé avant la demande d’annulation.")
            window.interrupt_queries()
            settle()
            expect(window.case.path == case_path and window.view == saved_view, "L’analyse précédente n’a pas été restaurée.")
            expect(window.tree.isColumnHidden(5) and window.case.bookmarks() == {1: "À conserver"}, "Colonnes ou notes non restaurées.")
        check("cancel_import_and_restore_previous_workspace", cancelled_import)

        def boolean_rules_and_sort():
            expression = FilterGroup("and", [FilterGroup("or", [Filter(1, "eq", "élève"), Filter(1, "eq", "bob")]),
                                            FilterGroup("not", [Filter(2, "eq", "2")])])
            rule = FormatRule(FilterGroup("and", [Filter(2, "gt", "9007199254740992")]),
                              "#ffc7ce", "#111111", True, 2)
            view = View(expression=expression, format_rules=[rule], sorts=[SortKey(0), SortKey(2, "number", True)])
            window.set_view(view)
            settle()
            expect([r[0] for r in window.case.page(window.view)] == [2, 1], "Expression imbriquée ou tri multi-colonnes incorrect.")
            expect(window.model.data(window.model.index(0, 3), Qt.BackgroundRole).name() == "#ffc7ce", "Coloration de cellule non appliquée.")
            expect(window.model.data(window.model.index(0, 3), Qt.FontRole).bold(), "Mise en gras non appliquée.")
            window.tree.clearSelection()
            rectangle = window.tree.visualRect(window.model.index(0, 3))
            pixel = window.tree.viewport().grab().toImage().pixelColor(rectangle.right() - 8, rectangle.bottom() - 4)
            expect(pixel.name() == "#ffc7ce", "La palette de la grille masque la couleur de cellule à l’écran.")
            window.case.save_view("avancé", window.view)
            expect(window.case.views()["avancé"] == window.view, "Vue avancée non persistée.")
        check("boolean_expression_multisort_and_cell_formatting", boolean_rules_and_sort)

        def tabs_and_session():
            from fireblade.core.engine import Case
            old_connection = window.case.db
            previous_view = window.view
            second_csv = directory / "deuxième.csv"
            second_csv.write_text("Name\nDeuxième fichier\n", encoding="utf-8")
            second_path = directory / "deuxième.timeline"
            other = Case(second_path)
            other.import_csv(second_csv)
            other.close()
            window.load_case(second_path)
            settle()
            expect(window.tabs.count() == 2 and window.case.path == second_path, "Deuxième onglet absent.")
            released = False
            try:
                old_connection.execute("SELECT 1")
            except duckdb.ConnectionException:
                released = True
            expect(released, "La connexion de l’onglet inactif consomme encore des ressources.")
            window.switch_document(0)
            settle()
            expect(window.view == previous_view and window.tree.isColumnHidden(5), "Vue ou colonnes perdues au changement d’onglet.")
            session = directory / "enquête.fireblade-session"
            window.save_session(session)
            window.close_document(0)
            settle()
            window.load_session(session)
            settle()
            expect(window.tabs.count() == 2 and window.case.path == case_path and window.view == previous_view,
                   "La session n’a pas restauré ses onglets et filtres.")
        check("tabs_release_inactive_connection_and_session_restore", tabs_and_session)

        def excel_roundtrip():
            from openpyxl import Workbook, load_workbook
            excel_source = directory / "preuve été.xlsx"
            workbook = Workbook()
            workbook.active.append(["Ignore"])
            sheet = workbook.create_sheet("Événements")
            sheet.append(["ID", "Details"])
            sheet.append(["0001", "=1+1"])
            sheet.append(["9007199254740993", "Ligne A\nLigne B"])
            workbook.save(excel_source)
            workbook.close()
            before = hashlib.sha256(excel_source.read_bytes()).hexdigest()
            import_with_dialog(excel_source, directory / "excel.timeline", "Événements")
            settle()
            expect(window.case.page(View()) == [(1, "0001", "=1+1"), (2, "9007199254740993", "Ligne A\nLigne B")], "Import Excel altéré.")
            expect(window.case.info["sha256"] == before, "La provenance Excel ne référence pas le classeur source.")
            output = directory / "résultats.xlsx"
            window.case.export(output, View())
            result = load_workbook(output, read_only=True, data_only=False)
            try:
                rows = list(result.active.iter_rows())
                expect(rows[1][1].data_type == "s" and rows[1][1].value == "=1+1", "Une formule exportée n’est pas du texte littéral.")
                expect(rows[2][0].value == "9007199254740993", "Grand entier altéré dans l’export Excel.")
            finally:
                result.close()
        check("xlsx_sheet_import_export_and_literal_formulas", excel_roundtrip)

        def shared_rules_across_columns():
            book = RuleBook(["User", "Absent"], [
                FormatRule(FilterGroup("and", [Filter(0, "eq", "élève")]), "#ffc7ce", target=0),
                FormatRule(FilterGroup("and", [Filter(1, "eq", "x")]), "#aabbcc")])
            window.rule_store.replace(book)
            window.reload_global_rules()
            settle()
            expect("non applicable" in window.rules_summary.text(), "Les règles inapplicables ne sont pas signalées.")
            rearranged = directory / "colonnes réordonnées.csv"
            rearranged.write_text("User,RecordID,Artifact\nélève,0001,MFT\nbob,9007199254740993,EVTX\nélève,2,MFT\n", encoding="utf-8")
            import_with_dialog(rearranged, directory / "colonnes.timeline")
            settle()
            expect(window.model.data(window.model.index(0, 1), Qt.BackgroundRole).name() == "#ffc7ce", "La règle globale ne suit pas le nom de colonne.")
            window.set_view(View(format_rules=[FormatRule(FilterGroup("and", [Filter(0, "eq", "élève")]), "#aabbcc", target=0)]))
            settle()
            expect(window.model.data(window.model.index(0, 1), Qt.BackgroundRole).name() == "#aabbcc", "La règle locale doit être prioritaire.")
            window.set_view(View(global_rules_enabled=False))
            settle()
            expect(window.model.data(window.model.index(0, 1), Qt.BackgroundRole) is None, "La désactivation globale par vue ne fonctionne pas.")
            expect(RuleStore(window.rule_store.path).load() == book, "Les règles globales n’ont pas été persistées.")
            before = window.rule_store.path.read_bytes()
            denied = False
            with window.rule_store.path.open("rb"):
                try:
                    window.rule_store.save(RuleBook())
                except PermissionError:
                    denied = True
            expect(denied and window.rule_store.path.read_bytes() == before, "Le verrouillage Windows des règles n’a pas préservé la configuration.")
            expect(not list(directory.glob("*.partial")), "Un fichier de règles temporaire reste après l’échec.")
        check("shared_rules_column_names_priority_persistence_and_windows_lock", shared_rules_across_columns)

        def selection_notes_and_export():
            selection = window.tree.selectionModel()
            for number in (0, 1):
                selection.select(window.model.index(number, 0), QItemSelectionModel.Select | QItemSelectionModel.Rows)
            window.annotate_selection("mark")
            settle()
            expect(window.case.bookmarks() == {1: "", 2: ""}, "La sélection n’a pas été marquée exactement.")
            selection = window.tree.selectionModel()
            for number in (0, 1):
                selection.select(window.model.index(number, 0), QItemSelectionModel.Select | QItemSelectionModel.Rows)
            window.annotate_selection("append", "À examiner\n=1+1")
            settle()
            expect(window.case.bookmarks() == {1: "À examiner\n=1+1", 2: "À examiner\n=1+1"}, "L’annotation collective est incorrecte.")
            expect(window.model.notes == window.case.bookmarks(), "Les notes de la page visible n’ont pas été actualisées.")
            exported = directory / "signets et notes.csv"
            window.case.export(exported, View(bookmarks_only=True), include_annotations=True)
            with exported.open(encoding="utf-8-sig", newline="") as incoming:
                rows = list(csv.reader(incoming))
            expect(len(rows) == 3 and rows[0][-3:] == ["Fireblade_SourceRecord", "Fireblade_Bookmarked", "Fireblade_Note"], "L’export annoté n’a pas les bonnes colonnes ou lignes.")
            expect(rows[1][-1] == "À examiner\n=1+1" and rows[2][1] == "9007199254740993", "Notes ou identifiants altérés dans l’export annoté.")
        check("bulk_selection_notes_and_annotated_export", selection_notes_and_export)

        report["memory"] = windows_memory()
        report["status"] = "passed"
    except Exception as exc:
        report["status"] = "failed"
        report["error"] = str(exc) or type(exc).__name__
        report["ui_errors"] = list(errors)
        report["current_case"] = str(window.case.path) if window is not None and window.case is not None else None
    finally:
        if window is not None:
            try:
                if window.jobs or window.pool.activeThreadCount() or (window.model and window.model.jobs):
                    window.interrupt_queries()
                    settle(timeout=10)
                if not window.close():
                    raise RuntimeError("La fenêtre de contrôle n’a pas pu être fermée.")
            except Exception as exc:
                report["status"] = "failed"
                report["cleanup_error"] = str(exc)
        for path in readonly:
            try:
                path.chmod(stat.S_IREAD | stat.S_IWRITE)
            except OSError:
                pass
        if work is not None:
            try:
                work.cleanup()
            except OSError as exc:
                report["status"] = "failed"
                report["cleanup_error"] = str(exc)
    report["finished_utc"] = datetime.now(timezone.utc).isoformat()
    report_path.parent.mkdir(parents=True, exist_ok=True)
    with report_path.open("x", encoding="utf-8") as output:
        json.dump(report, output, indent=2, ensure_ascii=False)
    return 0 if report["status"] == "passed" else 1
