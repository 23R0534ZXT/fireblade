import csv
import hashlib
import json
import os
from pathlib import Path
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import duckdb
import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QCoreApplication, QEvent, Qt
from PySide6.QtWidgets import QApplication

from fireblade.ui.advanced_dialogs import ExpressionEditor, RulesDialog, SortDialog
from fireblade.ui.app import Window
from fireblade.core.engine import Case
from fireblade.core.errors import OperationCancelled
from fireblade.core.models import Filter, FilterGroup, FormatRule, SortKey, View
from fireblade.core.profiles import detect_profile, profile_view
from fireblade.ui.workspace import read_session


@pytest.fixture
def case(tmp_path):
    source = tmp_path / "evidence.csv"
    source.write_text('User,EventID,Details\nalice,4624,"{\"\"ok\"\": true}"\nbob,4625,failed\nalice,1,other\n,4625,unknown\n', encoding="utf-8")
    case = Case(tmp_path / "evidence.timeline")
    case.import_csv(source)
    yield case
    case.close()


def test_nested_boolean_conditions_share_export_and_group_semantics(case, tmp_path):
    expression = FilterGroup("and", [
        FilterGroup("or", [Filter(0, "eq", "alice"), Filter(0, "eq", "bob")]),
        FilterGroup("not", [Filter(1, "eq", "1")])])
    view = View(expression=expression, groups=[0])
    assert case.count(view) == 2
    assert case.page(view) == [("alice", 1), ("bob", 1)]
    assert len(case.page(view, ("bob",))) == 1
    target = tmp_path / "result.csv"
    case.export(target, view)
    with target.open(encoding="utf-8-sig", newline="") as incoming:
        assert [r[0] for r in list(csv.reader(incoming))[1:]] == ["alice", "bob"]
    assert case.count(View(expression=FilterGroup("not", [Filter(0, "eq", "alice")]))) == 2
    assert case.count(View(expression=FilterGroup("or", []))) == 0
    injection = FilterGroup("or", [Filter(0, "eq", "' OR 1=1 --"), Filter(1, "eq", "1")])
    assert case.count(View(expression=injection)) == 1
    with pytest.raises(ValueError, match="NON"):
        case.count(View(expression=FilterGroup("not", [])))


def test_legacy_views_and_new_views_roundtrip(case):
    assert View.from_json('{"search":"alice","filters":[]}') == View(search="alice")
    view = View(expression=FilterGroup("or", [Filter(1, "eq", "4625"), Filter(0, "empty")]),
                format_rules=[FormatRule(FilterGroup("and", [Filter(1, "ge", "4625")]), "#ffc7ce", "#111111", True, 1)],
                sorts=[SortKey(0), SortKey(1, "number", True)])
    assert View.from_json(view.to_json()) == view
    case.save_view("advanced", view)
    assert case.views()["advanced"] == view


def test_multi_column_sort_matches_pages_and_export(case, tmp_path):
    view = View(sorts=[SortKey(0, "text", True), SortKey(1, "number", True)])
    assert [row[0] for row in case.page(view)] == [2, 1, 3, 4]
    target = tmp_path / "sorted.csv"
    case.export(target, view)
    with target.open(encoding="utf-8-sig", newline="") as incoming:
        assert [r[1] for r in list(csv.reader(incoming))[1:]] == ["4625", "4624", "1", "4625"]
    with pytest.raises(ValueError, match="distinctes"):
        case.page(View(sorts=[SortKey(0), SortKey(0)]))


def test_condition_budget_and_rule_validation(case):
    expression = FilterGroup("and", [Filter(0, "eq", "alice")] * 1025)
    with pytest.raises(ValueError, match="1024"):
        case.count(View(expression=expression))
    view = View(format_rules=[FormatRule(background="invalid")])
    with pytest.raises(ValueError, match="Couleur"):
        case.validate_view(view)
    nested = Filter(0, "eq", "alice")
    for _ in range(26):
        nested = FilterGroup("and", [nested])
    with pytest.raises(ValueError, match="24"):
        case.count(View(expression=nested))


def test_styles_match_sql_and_only_requested_rows(case):
    rules = [FormatRule(FilterGroup("and", [Filter(1, "ge", "4625")]), target=1),
             FormatRule(FilterGroup("or", [Filter(0, "eq", "alice"), Filter(0, "empty")])),
             FormatRule(enabled=False)]
    view = View(format_rules=rules)
    visible_rows = case.page(view)[:2]
    assert case.page_styles(view, visible_rows) == {1: (1,), 2: (0,)}
    # Formatting must describe the values actually displayed, not a second read
    # of a newer table version, and must not scan the full event table.
    case.db.execute("DROP TABLE events")
    assert case.page_styles(view, visible_rows) == {1: (1,), 2: (0,)}
    assert case.page_styles(view, []) == {}


def workbook_source(path):
    book = Workbook()
    book.active.title = "Sans intérêt"
    book.active.append(["Other"])
    sheet = book.create_sheet("Événements")
    sheet.append(["ID", "Details"])
    sheet.append(["0001", "=1+1"])
    sheet.append(["9007199254740993", "line A\nline B"])
    book.save(path)
    book.close()


def test_xlsx_import_keeps_formula_text_sheet_and_source_hash(tmp_path):
    source = tmp_path / "preuve été.xlsx"
    workbook_source(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    case = Case(tmp_path / "excel.timeline")
    try:
        case.import_excel(source, "Événements")
        assert case.page(View()) == [(1, "0001", "=1+1"), (2, "9007199254740993", "line A\nline B")]
        assert case.info["sha256"] == before
        assert case.info["path"] == str(source.resolve())
        assert case.info["sheet"] == "Événements"
        assert hashlib.sha256(source.read_bytes()).hexdigest() == before
        assert not list(tmp_path.glob("tmp*.csv"))
    finally:
        case.close()


def test_xlsx_export_preserves_text_instead_of_formula_execution(case, tmp_path):
    case.db.execute("UPDATE events SET c2='=1+1' WHERE __row_id=1")
    target = tmp_path / "result.xlsx"
    case.export(target, View())
    workbook = load_workbook(target, read_only=True, data_only=False)
    try:
        rows = list(workbook.active.iter_rows())
        assert rows[1][2].value == "=1+1" and rows[1][2].data_type == "s"
        assert rows[1][1].value == "4624" and rows[1][1].data_type == "s"
        assert len(rows) == 5
    finally:
        workbook.close()
    assert not target.with_name(target.name + ".partial").exists()


def test_xlsx_cell_limit_does_not_truncate_or_replace_old_export(case, tmp_path):
    case.db.execute("UPDATE events SET c2=? WHERE __row_id=1", ["x" * 32768])
    target = tmp_path / "result.xlsx"
    target.write_bytes(b"previous file")
    with pytest.raises(ValueError, match="32767"):
        case.export(target, View())
    assert target.read_bytes() == b"previous file"
    assert not target.with_name(target.name + ".partial").exists()


def test_cancel_xlsx_import_rolls_back_and_can_retry(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook_source(source)
    case = Case(tmp_path / "case.timeline")
    try:
        def cancel(message):
            if message.startswith("Écriture"):
                case.interrupt()
        with pytest.raises(OperationCancelled):
            case.import_excel(source, "Événements", cancel)
        assert case.headers == []
        assert not list(tmp_path.glob("tmp*.csv"))
        case.import_excel(source, "Événements")
        assert case.count(View()) == 2
    finally:
        case.close()


def test_profiles_require_signature_and_keep_unknown_columns():
    assert detect_profile(["EventId"]) is None
    headers = ["Channel", "EventId", "TimeCreated", "Provider", "CustomField"]
    profile = detect_profile(headers)
    assert profile.name == "EvtxECmd"
    view, order = profile_view(profile, headers)
    assert view.sort_column == 2 and view.sort_type == "timestamp"
    assert sorted(order) == list(range(5))


def settle(app, window):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        app.processEvents()
        if not window.jobs and not window.pool.activeThreadCount() and (window.model is None or not window.model.jobs):
            return
        time.sleep(.005)
    pytest.fail("Background work timed out")


def test_tabs_release_inactive_database_and_restore_views_and_session(case, tmp_path):
    app = QApplication.instance() or QApplication([])
    window = Window()
    errors = []
    window.error = lambda e: errors.append(str(e))
    window.attach(case)
    window.show()
    settle(app, window)
    first = case.path
    window.set_view(View(search="alice"))
    settle(app, window)
    window.tree.setColumnHidden(2, True)
    source = tmp_path / "second.csv"
    source.write_text("Name\nsecond\n")
    second = Case(tmp_path / "second.timeline")
    second.import_csv(source)
    second.close()
    window.load_case(second.path)
    settle(app, window)
    assert window.tabs.count() == 2
    with pytest.raises(duckdb.ConnectionException):
        case.db.execute("SELECT 1")
    # A closed inactive Windows file can be renamed.
    moved = first.with_suffix(".moved")
    first.rename(moved)
    moved.rename(first)
    window.tabs.setCurrentIndex(0)
    settle(app, window)
    assert window.case.path == first and window.view.search == "alice"
    assert window.tree.isColumnHidden(2)
    assert window.case.count(window.view) == 2
    session = tmp_path / "enquête.fireblade-session"
    window.save_session(session)
    assert session.is_file()
    window.close_document(0)
    settle(app, window)
    assert window.tabs.count() == 1 and window.case.path == second.path
    window.load_session(session)
    settle(app, window)
    assert window.tabs.count() == 2 and window.case.path == first
    assert window.view.search == "alice" and window.tree.isColumnHidden(2)
    assert not errors
    assert window.close()


def test_invalid_session_preserves_current_analysis(case, tmp_path):
    app = QApplication.instance() or QApplication([])
    window = Window()
    window.attach(case)
    settle(app, window)
    errors = []
    window.error = lambda e: errors.append(str(e))
    session = tmp_path / "test.fireblade-session"
    window.save_session(session)
    data = json.loads(session.read_text(encoding="utf-8"))
    data["documents"][0]["source"]["sha256"] = "wrong source"
    session.write_text(json.dumps(data), encoding="utf-8")
    window.load_session(session)
    assert errors and "remplacée" in errors[0]
    assert window.case is case and case.count(View()) == 4
    assert window.close()


def test_editors_roundtrip_nested_expression_and_format_priorities(case):
    app = QApplication.instance() or QApplication([])
    expression = FilterGroup("or", [Filter(0, "eq", "alice"), FilterGroup("not", [Filter(1, "gt", "1")])])
    editor = ExpressionEditor(case.headers, expression)
    assert editor.value() == expression
    rules = RulesDialog(case.headers, [FormatRule(expression, "#ff0000"), FormatRule(background="#00ff00")])
    rules.move(1)
    rules.accept()
    assert rules.rules[1].condition == expression
    assert rules.rules[1].background == "#ff0000"
    sorts = SortDialog(case.headers, [SortKey(0), SortKey(1, "number", True)])
    sorts.move(-1)
    assert sorts.values() == [SortKey(1, "number", True), SortKey(0)]
    for widget in (editor, rules, sorts):
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


def test_format_rule_precedes_bookmark_and_targets_only_one_cell(case):
    app = QApplication.instance() or QApplication([])
    case.bookmark(1, "investigate")
    window = Window()
    window.attach(case)
    window.show()
    settle(app, window)
    view = View(format_rules=[FormatRule(FilterGroup("and", [Filter(0, "eq", "alice")]), "#ff0000", "#0000ff", True, 1)])
    window.set_view(view)
    settle(app, window)
    assert window.model.data(window.model.index(0, 2), Qt.BackgroundRole).name() == "#ff0000"
    assert window.model.data(window.model.index(0, 2), Qt.FontRole).bold()
    assert window.model.data(window.model.index(0, 1), Qt.BackgroundRole).name() == "#fff4c2"
    window.tree.clearSelection()
    rectangle = window.tree.visualRect(window.model.index(0, 2))
    pixel = window.tree.viewport().grab().toImage().pixelColor(rectangle.right() - 8, rectangle.bottom() - 4)
    assert pixel.name() == "#ff0000"
    assert window.close()


def test_rules_dialog_click_creation_replaces_visible_editor_and_saves(case, monkeypatch):
    from PySide6.QtGui import QColor
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QColorDialog, QPushButton, QDialogButtonBox, QDialog
    app = QApplication.instance() or QApplication([])
    dialog = RulesDialog(case.headers, [])
    dialog.show()
    app.processEvents()
    def click(text):
        button = next(b for b in dialog.findChildren(QPushButton) if b.text() == text and b.isVisible() and b.isEnabled())
        QTest.mouseClick(button, Qt.LeftButton)
        app.processEvents()
    try:
        click('Ajouter')
        # processEvents deliberately does not force DeferredDelete delivery:
        # these controls must already be retired while a modal loop is active.
        assert sum(e.isVisible() for e in dialog.findChildren(ExpressionEditor)) == 1
        item = dialog.editor.root.child(0)
        tree = dialog.editor.tree
        tree.itemWidget(item, 1).setCurrentIndex(tree.itemWidget(item, 1).findData('eq'))
        QTest.keyClicks(tree.itemWidget(item, 2), 'alice')
        monkeypatch.setattr(QColorDialog, 'getColor', lambda *a, **k: QColor('#ff0000'))
        click('Fond')
        QTest.mouseClick(dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Ok), Qt.LeftButton)
        assert dialog.result() == QDialog.DialogCode.Accepted
        assert dialog.rules[0].condition == FilterGroup('and', [Filter(0, 'eq', 'alice')])
        assert dialog.rules[0].background == '#ff0000'
        view = View(format_rules=dialog.rules)
        assert case.page_styles(view, case.page(view)) == {1: (0,), 2: (), 3: (0,), 4: ()}
    finally:
        dialog.close()
        dialog.deleteLater()
