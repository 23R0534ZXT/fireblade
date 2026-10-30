import csv
import hashlib
import time

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QItemSelectionModel, Qt

from fireblade.ui.app import Window
from fireblade.core.engine import Case
from fireblade.core.errors import OperationCancelled
from fireblade.core.models import Filter, FilterGroup, FormatRule, View
from fireblade.core.rule_store import RuleBook, RuleStore, read_book
from fireblade.core.settings import Settings


def rule(column=0, value="alice", color="#ffcccc", target=None):
    return FormatRule(FilterGroup("and", [Filter(column, "eq", value)]), color, target=target)


@pytest.fixture
def case(tmp_path):
    source = tmp_path / "evidence.csv"
    source.write_text("User,ID,Details\nalice,0001,first\nbob,9007199254740993,second\nalice,2,third\n", encoding="utf-8")
    case = Case(tmp_path / "evidence.timeline")
    case.import_csv(source)
    yield case
    case.close()


def settle(app, window):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        app.processEvents()
        if not window.jobs and not window.pool.activeThreadCount() and (window.model is None or not window.model.jobs):
            return
        time.sleep(.005)
    pytest.fail("Background work timed out")


def test_global_rules_remap_by_name_and_skip_whole_incomplete_rule(case):
    book = RuleBook(["ID", "USER", "Absent"], [rule(1, target=1), rule(2)])
    adapted, skipped = book.adapt(case.headers)
    assert adapted[0].condition.children[0].column == 0
    assert adapted[0].target == 0
    assert len(skipped) == 1 and "Absent" in skipped[0]
    assert case.page_styles(View(), case.page(View()), adapted) == {1: (0,), 2: (), 3: (0,)}
    nested = FormatRule(FilterGroup("or", [Filter(1, "eq", "alice"), Filter(2, "eq", "x")]))
    assert RuleBook(book.headers, [nested]).adapt(case.headers)[0] == []
    assert RuleBook(["User", "Gone"], [rule(0, target=1)]).adapt(case.headers)[0] == []


def test_ambiguous_column_case_is_not_guessed():
    book = RuleBook(["USER"], [rule()])
    assert book.adapt(["User", "user"])[0] == []
    assert RuleBook(["User"], [rule()]).adapt(["User", "user"])[0][0].condition.children[0].column == 0


def test_global_rules_have_separate_limit_and_cannot_expand_local_saved_views(case):
    local = [rule(target=0)] * 32
    shared = [rule(color="#aabbcc", target=0)] * 32
    view = View(format_rules=local)
    case.validate_view(view)
    assert len(case.page_styles(view, case.page(view), shared)[1]) == 64
    assert len(View.from_json(view.to_json()).format_rules) == 32
    assert case.page_styles(View(global_rules_enabled=False), case.page(View()), shared) == {}


@pytest.mark.parametrize("condition", [Filter(0, "regex", "["), Filter(0, "ge", "not a number"), Filter(0, "after", "invalid date")])
def test_bad_rule_operand_never_replaces_existing_file(tmp_path, condition):
    path = tmp_path / "rules.json"
    store = RuleStore(path)
    store.save(RuleBook(["User"], [rule()]))
    original = path.read_bytes()
    invalid = RuleBook(["User"], [FormatRule(FilterGroup("and", [condition]))])
    with pytest.raises(ValueError):
        store.save(invalid)
    assert path.read_bytes() == original
    assert not list(tmp_path.glob("*.partial"))


def test_store_roundtrip_stale_editor_and_failed_atomic_write(tmp_path, monkeypatch):
    import fireblade.core.rule_store as storage
    path = tmp_path / "rules.json"
    first, other = RuleStore(path), RuleStore(path)
    book = RuleBook(["User"], [rule()])
    first.save(book)
    assert other.load() == book
    first.save(RuleBook(["User"], [rule(color="#aabbcc")]))
    with pytest.raises(ValueError, match="changé"):
        other.save(book)
    assert read_book(path).rules[0].background == "#aabbcc"
    old = path.read_bytes()
    def denied(*args):
        raise PermissionError("locked destination")
    monkeypatch.setattr(storage.os, "replace", denied)
    with pytest.raises(PermissionError):
        first.save(book)
    assert path.read_bytes() == old
    assert not list(tmp_path.glob("*.partial"))


def test_global_rules_invalid_file_is_recoverable_without_changing_analysis(qt_application, case, tmp_path):
    path = tmp_path / "global.json"
    path.write_bytes(b"invalid JSON")
    window = Window(rule_path=path)
    errors = []
    window.error = lambda error: errors.append(str(error))
    window.attach(case)
    settle(qt_application, window)
    assert window.global_rules_error and "invalide" in window.rules_summary.text()
    assert case.count(window.view) == 3 and path.read_bytes() == b"invalid JSON"
    valid = tmp_path / "import.json"
    valid.write_bytes(RuleBook(["User"], [rule()]).to_bytes())
    window.import_global_rules(valid)
    settle(qt_application, window)
    assert window.global_rules_error is None
    assert window.model.data(window.model.index(0, 1), Qt.BackgroundRole).name() == "#ffcccc"
    assert not errors
    assert window.close()


def test_global_local_precedence_and_per_view_disable_persist(qt_application, case, tmp_path):
    path = tmp_path / "global.json"
    RuleStore(path).save(RuleBook(["User"], [rule(target=0)]))
    window = Window(rule_path=path)
    window.attach(case)
    window.show()
    settle(qt_application, window)
    assert not window.view.format_rules  # Shared rules never leak into saved local views.
    assert window.model.data(window.model.index(0, 1), Qt.BackgroundRole).name() == "#ffcccc"
    window.set_view(View(format_rules=[rule(color="#aabbcc", target=0)]))
    settle(qt_application, window)
    assert window.model.data(window.model.index(0, 1), Qt.BackgroundRole).name() == "#aabbcc"
    window.set_view(View(global_rules_enabled=False))
    settle(qt_application, window)
    assert window.model.data(window.model.index(0, 1), Qt.BackgroundRole) is None
    assert not window.global_rules_action.isChecked()
    case.save_view("without globals", window.view)
    assert not case.views()["without globals"].global_rules_enabled
    assert View.from_json('{"filters":[]}').global_rules_enabled
    assert window.close()
    assert RuleStore(path).load().rules == [rule(target=0)]


def test_rule_editor_keeps_fields_for_other_artifact_types(qt_application, case, tmp_path, monkeypatch):
    import fireblade.ui.shared_rules as ui
    path = tmp_path / "global.json"
    book = RuleBook(["DifferentArtifactField"], [rule(value="critical")])
    RuleStore(path).save(book)
    captured = []
    def accept(dialog):
        captured.extend(dialog.headers)
        dialog.add()
        # Select the current CSV's User column; keep the foreign rule unchanged.
        dialog.editor.tree.itemWidget(dialog.editor.root.child(0), 0).setCurrentIndex(dialog.headers.index("User"))
        dialog.editor.tree.itemWidget(dialog.editor.root.child(0), 2).setText("alice")
        dialog.accept()
        return dialog.result()
    monkeypatch.setattr(ui.RulesDialog, "exec", accept)
    window = Window(rule_path=path)
    window.attach(case)
    settle(qt_application, window)
    window.edit_global_rules()
    settle(qt_application, window)
    saved = read_book(path)
    assert "DifferentArtifactField" in captured and "User" in captured
    assert saved.rules[0] == book.rules[0]
    assert len(saved.rules) == 2 and len(saved.adapt(case.headers)[0]) == 1
    assert window.close()


def test_bulk_bookmark_modes_preserve_notes_and_original_source(case):
    source = case.info["path"]
    from pathlib import Path
    before = hashlib.sha256(Path(source).read_bytes()).hexdigest()
    case.bookmark(1, "Original note")
    assert case.bookmark_many([1, 2, 2], mode="mark") == 2
    assert case.bookmarks() == {1: "Original note", 2: ""}
    case.bookmark_many([1, 2], "Événement à examiner", "append")
    assert case.bookmarks() == {1: "Original note\nÉvénement à examiner", 2: "Événement à examiner"}
    case.bookmark_many([2], "' OR 1=1 --", "replace")
    assert case.bookmarks([2]) == {2: "' OR 1=1 --"}
    assert case.bookmarks([]) == {}
    case.bookmark_many([1], mode="remove")
    assert case.bookmarks() == {2: "' OR 1=1 --"}
    assert case.count(View(bookmarks_only=True)) == 1
    with pytest.raises(ValueError):
        case.bookmark_many([2, 999], mode="remove")
    assert case.bookmarks() == {2: "' OR 1=1 --"}
    assert hashlib.sha256(Path(source).read_bytes()).hexdigest() == before


def test_bulk_bookmark_cancel_rolls_back_entire_selection(case, monkeypatch):
    case.bookmark(1, "Keep")
    revision = case._bookmark_revision
    original = case.check_cancelled
    def before_commit():
        if case._finalizing_transaction:
            case.interrupt()
        original()
    with monkeypatch.context() as patch:
        patch.setattr(case, "check_cancelled", before_commit)
        with pytest.raises(OperationCancelled):
            case.bookmark_many([1, 2, 3], "Must not persist", "replace")
    assert case.bookmarks() == {1: "Keep"}
    assert case._bookmark_revision == revision
    case.bookmark_many([2, 3], mode="mark")
    assert set(case.bookmarks()) == {1, 2, 3}


def test_note_export_preserves_filter_order_source_rows_and_text(case, tmp_path):
    case.bookmark_many([1, 2, 3], "=1+1\nÉvidence", "replace")
    view = View(filters=[Filter(0, "eq", "alice")], bookmarks_only=True, sort_column=1, sort_type="number", descending=True)
    csv_path = tmp_path / "notes.csv"
    case.export(csv_path, view, include_annotations=True)
    with csv_path.open(encoding="utf-8-sig", newline="") as incoming:
        rows = list(csv.reader(incoming))
    assert rows[0][-3:] == ["Fireblade_SourceRecord", "Fireblade_Bookmarked", "Fireblade_Note"]
    assert [r[-3] for r in rows[1:]] == ["3", "1"]
    assert all(r[-1] == "=1+1\nÉvidence" for r in rows[1:])
    xlsx_path = tmp_path / "notes.xlsx"
    case.export(xlsx_path, view, include_annotations=True)
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        rows = list(workbook.active.iter_rows())
        assert rows[1][-1].data_type == "s" and rows[1][-1].value == "=1+1\nÉvidence"
        assert rows[2][1].value == "0001"
    finally:
        workbook.close()
    with pytest.raises(ValueError):
        case.export(tmp_path / "groups.csv", View(groups=[0]), grouped=True, include_annotations=True)


def test_annotation_column_names_do_not_hide_existing_evidence_columns(tmp_path):
    source = tmp_path / "source.csv"
    source.write_text("User,Fireblade_Note\nalice,original value\n")
    case = Case(tmp_path / "source.timeline")
    case.import_csv(source)
    try:
        case.bookmark(1, "analyst note")
        target = tmp_path / "annotated.csv"
        case.export(target, View(), include_annotations=True)
        with target.open(encoding="utf-8-sig", newline="") as incoming:
            rows = list(csv.reader(incoming))
        assert rows[0][1] == "Fireblade_Note" and rows[0][-1] == "Fireblade_Note_2"
        assert rows[1][1] == "original value" and rows[1][-1] == "analyst note"
    finally:
        case.close()


def test_selection_ignores_group_headers_and_exports_notes(qt_application, case, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    window = Window()
    window.attach(case)
    settle(qt_application, window)
    window.set_view(View(groups=[0]))
    settle(qt_application, window)
    parent = window.model.index(0, 0)
    window.model.fetchMore(parent)
    settle(qt_application, window)
    selection = window.tree.selectionModel()
    selection.select(parent, QItemSelectionModel.Select | QItemSelectionModel.Rows)
    selection.select(window.model.index(0, 0, parent), QItemSelectionModel.Select | QItemSelectionModel.Rows)
    window.annotate_selection("append", "Review")
    settle(qt_application, window)
    assert len(case.bookmarks()) == 1 and list(case.bookmarks().values()) == ["Review"]
    target = tmp_path / "notes.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(target), ""))
    window.export_annotations()
    settle(qt_application, window)
    with target.open(encoding="utf-8-sig", newline="") as incoming:
        assert len(list(csv.reader(incoming))) == 2
    assert window.close()


def test_annotation_cache_is_bounded_by_loaded_pages(qt_application, tmp_path):
    source = tmp_path / "many.csv"
    source.write_text("ID,User\n" + "".join(f"{i},alice\n" for i in range(150)))
    case = Case(tmp_path / "many.timeline", settings=Settings(page_size=25, max_loaded_rows=50))
    case.import_csv(source)
    case.bookmark_many(list(range(1, 151)), "Long note " * 300, "replace")
    window = Window()
    window.attach(case)
    settle(qt_application, window)
    assert len(window.model.notes) == 25
    window.more()
    settle(qt_application, window)
    assert len(window.model.notes) == 50
    window.more()
    settle(qt_application, window)
    assert set(window.model.notes) == set(range(51, 76))
    window.set_view(View(groups=[1]))
    settle(qt_application, window)
    assert window.model.notes == {}
    index = window.model.index(0, 0)
    window.model.fetchMore(index)
    settle(qt_application, window)
    assert len(window.model.notes) == 25
    window.model.unload(index)
    assert window.model.notes == {}
    assert len(case.bookmarks()) == 150
    assert window.close()
