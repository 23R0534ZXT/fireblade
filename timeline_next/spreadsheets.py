"""Streaming XLSX support. Formulas are read and exported as literal text."""
import csv
from datetime import date, datetime, time
from pathlib import Path
import tempfile
from contextlib import contextmanager

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


def sheet_names(source):
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


@contextmanager
def csv_from_sheet(source, sheet, directory, check_cancelled):
    workbook, temporary = None, None
    try:
        workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Feuille Excel introuvable : {sheet}")
        worksheet = workbook[sheet]
        if worksheet.max_column and worksheet.max_column > 1024:
            raise ValueError("Une feuille de plus de 1024 colonnes n’est pas prise en charge")
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", suffix=".csv", dir=directory, delete=False) as output:
            temporary = Path(output.name)
            writer = csv.writer(output)
            for number, row in enumerate(worksheet.iter_rows(values_only=True)):
                if number % 256 == 0:
                    check_cancelled()
                values = [cell_text(value) for value in row]
                if number == 0:
                    if not values or not any(values):
                        raise ValueError("La première ligne doit contenir les noms des colonnes")
                writer.writerow(values)
        yield temporary
    finally:
        if workbook:
            workbook.close()
        if temporary:
            temporary.unlink(missing_ok=True)


def write_xlsx(path, headers, cursor, batch_size, check_cancelled):
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Résultats")
    worksheets = [worksheet]
    row_count = 0

    def append(values):
        nonlocal row_count, worksheet
        if row_count >= 1_048_576:
            worksheet = workbook.create_sheet(f"Résultats {len(worksheets) + 1}")
            worksheets.append(worksheet)
            row_count = 0
            append(headers)
        cells = []
        for value in values:
            if value is not None and len(str(value)) > 32767:
                raise ValueError("Une cellule dépasse la limite Excel de 32767 caractères ; exportez en CSV.")
            cell = WriteOnlyCell(worksheet, value=None if value is None else str(value))
            if value is not None:
                cell.data_type = "s"  # Never execute a CSV value beginning with '=' as a formula.
            cells.append(cell)
        worksheet.append(cells)
        row_count += 1

    try:
        append(headers)
        while rows := cursor.fetchmany(batch_size):
            check_cancelled()
            for row in rows:
                append(row)
        check_cancelled()
        workbook.save(path)
    finally:
        # openpyxl uses temporary sheet XML files, including on an aborted export.
        for sheet in worksheets:
            try:
                if not sheet.closed:
                    sheet.close()
            finally:
                writer = getattr(sheet, "_writer", None)
                if writer:
                    try:
                        writer.cleanup()
                    except FileNotFoundError:
                        pass
        workbook.close()
